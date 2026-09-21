import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_bright_aqua_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Bright Frutiger Aqua Palette (Inspired by bright tropical ocean artwork)
    BG_SKY_TOP = RGBColor(92, 211, 255)       # Bright sky blue
    CARD_BG = RGBColor(245, 252, 255)         # Crisp bright white-aqua card
    BORDER_CYAN = RGBColor(160, 225, 255)     # Soft cyan card border
    TEXT_TITLE = RGBColor(7, 42, 74)          # Deep clear navy title
    TEXT_SUB = RGBColor(2, 132, 199)          # Vibrant ocean blue subtext
    TEXT_HEAD = RGBColor(8, 47, 73)           # Dark crisp bullet heading
    TEXT_BODY = RGBColor(28, 61, 90)          # Clean readable body text
    TEXT_WHITE = RGBColor(255, 255, 255)

    # Bright Colorful Accents
    COLOR_GREEN = RGBColor(16, 185, 129)      # Bright Green (Strength)
    COLOR_AMBER = RGBColor(217, 119, 6)       # Warm Orange/Amber (Weakness)
    COLOR_BLUE = RGBColor(2, 132, 199)        # Bright Aqua Blue (Opportunity)
    COLOR_RED = RGBColor(220, 38, 38)         # Bright Red/Coral (Threat)
    COLOR_RAINBOW = RGBColor(255, 210, 0)     # Sunny gold accent

    blank_layout = prs.slide_layouts[6]

    def add_bright_background(slide):
        # Base background fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(230, 246, 255)
        bg.line.fill.background()

        # Top bright sky banner
        sky = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(2.2))
        sky.fill.solid()
        sky.fill.fore_color.rgb = BG_SKY_TOP
        sky.line.fill.background()

        # Rainbow accent line across top
        rb = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
        rb.fill.solid()
        rb.fill.fore_color.rgb = RGBColor(255, 77, 109)
        rb.line.fill.background()

        # Bubble shape on top-right
        b = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.2), Inches(0.35), Inches(1.6), Inches(1.6))
        b.fill.solid()
        b.fill.fore_color.rgb = RGBColor(255, 255, 255)
        b.line.color.rgb = RGBColor(160, 225, 255)
        b.line.width = Pt(1.5)

    slides_data = [
        # Slide 0: Title Cover
        {
            "is_cover": True,
            "badge": "WATER QUALITY STATION (WQS)",
            "title": "WQS Dashboard Module 1",
            "subtitle": "Product SWOT Analysis & Future Plan",
            "meta": "Automated Water Monitoring • In-House LoRa & ESP32 • 200 Ponds Live Control"
        },
        # Slide 1: Strength 1
        {
            "category": "STRENGTH",
            "category_color": COLOR_GREEN,
            "slide_num": "01 / 08",
            "title": "Auto Analyse & Actionable Data",
            "takeaway": "Automatic graphs and quick alerts for up to 200 ponds at the same time",
            "bullets": [
                ("Real-Time Data & Trend Graphs:", "Automatically records water data and draws trend graphs for up to 200 ponds at the same time without manual work."),
                ("Instant Warning When Values Go Wrong:", "Quickly detects abnormal readings like low oxygen or bad pH, warning the farm team before shrimp get hurt."),
                ("Easy View for Farm Managers:", "Managers can see the current water situation across all 200 ponds on one screen, without waiting for manual paper test logs.")
            ],
            "metric_box": ("200 PONDS", "Checked Together at the Same Time")
        },
        # Slide 2: Strength 2
        {
            "category": "STRENGTH",
            "category_color": COLOR_GREEN,
            "slide_num": "02 / 08",
            "title": "100% Internally Developed & Kaizen",
            "takeaway": "Full ownership of hardware and software, always improving",
            "bullets": [
                ("Built Completely In-House:", "I developed all the sensor electronics, circuit boards, code, and the web dashboard myself without hiring outside vendors."),
                ("No Outside Vendor Dependency:", "No expensive support contracts with 3rd-party companies, and no waiting weeks for outside technicians to come help."),
                ("Same-Day Repairs & Continuous Kaizen:", "If something breaks, we can fix it on the same day. We can also add new custom features anytime based on farm team feedback.")
            ],
            "metric_box": ("SAME DAY", "Fixes & Upgrades Done On-Site Immediately")
        },
        # Slide 3: Strength 3
        {
            "category": "STRENGTH",
            "category_color": COLOR_GREEN,
            "slide_num": "03 / 08",
            "title": "Cheap & Efficient",
            "takeaway": "Good performance without buying overpriced commercial equipment",
            "bullets": [
                ("Cost-Effective Microcontrollers & LoRa:", "Built with reliable, budget-friendly ESP32 chips and long-range wireless LoRa, which covers our large ponds easily."),
                ("No Overpowered Waste:", "We only use what our farm really needs. It costs only a small fraction of expensive commercial aquaculture packages."),
                ("Zero License Fees:", "Completely free from monthly or yearly software license fees, saving our company money every month.")
            ],
            "metric_box": ("$0 FEES", "No Yearly Per-Pond Subscription Fees")
        },
        # Slide 4: Weakness 1
        {
            "category": "WEAKNESS",
            "category_color": COLOR_AMBER,
            "slide_num": "04 / 08",
            "title": "Big Data",
            "takeaway": "Huge data coming in daily; need PostgreSQL instead of Excel",
            "bullets": [
                ("Huge Daily Rows:", "200 ponds sending data every 10 mins = 1,200 rows every hour, 28,800 rows every day, and ~10.5 Million rows each year!"),
                ("Complex ID Mapping Logic:", "Sensors only send general pond numbers (01.02.12). Our code must check the active shrimp list and attach the right cycle ID (e.g. 2010212.40)."),
                ("Need PostgreSQL Database Skills:", "Basic Excel or Google Sheets will freeze and crash with this much data. I need to master PostgreSQL database setup so graphs load fast without lag.")
            ],
            "metric_box": ("28,800 ROWS/DAY", "10.5M Rows/Year (Need PostgreSQL)")
        },
        # Slide 5: Weakness 2
        {
            "category": "WEAKNESS",
            "category_color": COLOR_AMBER,
            "slide_num": "05 / 08",
            "title": "Electronic Weakness",
            "takeaway": "Outdoor farm conditions are tough on electronic probes",
            "bullets": [
                ("Known Hardware Limitation:", "Even top brands like YSI advise checking probe measurements daily. Outdoor electronics always face heat, rain, and insects."),
                ("Algae Build-Up & Mineral Dirt:", "Algae and minerals in pond water stick to oxygen and pH probes, causing readings to drift if they are not cleaned regularly."),
                ("Weekly Maintenance Routine:", "We must set up a strict weekly routine for farm workers to wash probes and verify that sensor numbers are accurate.")
            ],
            "metric_box": ("WEEKLY", "Regular Probe Cleaning & Verification Schedule")
        },
        # Slide 6: Opportunity 1
        {
            "category": "OPPORTUNITY",
            "category_color": COLOR_BLUE,
            "slide_num": "06 / 08",
            "title": "Predictive Machine Learning",
            "takeaway": "Using data from 400 ponds per year to predict problems early",
            "bullets": [
                ("Big Data from 400 Ponds Per Year:", "Collect multi-cycle data across 400 ponds each year, combining harvest shrimp weight, water quality records, and daily feeding trends."),
                ("Forecast Problems Before They Happen:", "Train smart models to spot warning signs and predict oxygen drops or water sickness before shrimp show signs of stress."),
                ("Save Money on Feed Costs:", "Connect feeding amounts directly with dissolved oxygen and temperature to prevent overfeeding and cut expensive feed waste.")
            ],
            "metric_box": ("400 PONDS/YR", "Big Data to Train Smart Models")
        },
        # Slide 7: Opportunity 2
        {
            "category": "OPPORTUNITY",
            "category_color": COLOR_BLUE,
            "slide_num": "07 / 08",
            "title": "Automation",
            "takeaway": "Automatically control farm equipment based on water quality",
            "bullets": [
                ("Auto Paddlewheels & Auto Feeders:", "Automatically turn on paddlewheels when oxygen drops, and adjust auto-feeders to feed only when shrimp have good appetite."),
                ("Future: Underwater Cameras:", "Put cameras underwater to monitor uneaten feed pellets on the pond bottom and spot dead shrimp early."),
                ("Future: Automated Water Pumps:", "Automatically turn on water pumps for water exchange when water quality numbers (like pH or turbidity) get bad.")
            ],
            "metric_box": ("AUTO CONTROL", "Paddlewheels, Feeders & Pumps")
        },
        # Slide 8: Threat 1
        {
            "category": "THREAT",
            "category_color": COLOR_RED,
            "slide_num": "08 / 08",
            "title": "Harsh Farm Environment & Supply Chain",
            "takeaway": "Protecting electronics from outdoor weather and keeping spare parts",
            "bullets": [
                ("Tough Farm Weather Hazards:", "Outdoor electronics face heavy tropical thunderstorms, lightning power surges, high humidity, corrosion, and extreme heat."),
                ("Electronic Spare Parts Delays:", "If special sensor probes or chips run out of stock with suppliers, building or replacing units could face delays."),
                ("How We Protect It:", "Use IP67 waterproof boxes, lightning surge protectors, and always keep extra backup spare parts ready on the farm.")
            ],
            "metric_box": ("IP67 BOXES", "Waterproof & Surge Protected")
        }
    ]

    for item in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        add_bright_background(slide)

        if item.get("is_cover"):
            # Central Title Card
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.3), Inches(10.33), Inches(4.9))
            card.fill.solid()
            card.fill.fore_color.rgb = CARD_BG
            card.line.color.rgb = RGBColor(186, 230, 253)
            card.line.width = Pt(2)

            # Badge
            badge = slide.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.3), Inches(0.5))
            p_b = badge.text_frame.paragraphs[0]
            p_b.text = item["badge"]
            p_b.font.size = Pt(14)
            p_b.font.bold = True
            p_b.font.color.rgb = TEXT_SUB

            # Main Title
            title_box = slide.shapes.add_textbox(Inches(2.0), Inches(2.4), Inches(9.3), Inches(1.2))
            p_t = title_box.text_frame.paragraphs[0]
            p_t.text = item["title"]
            p_t.font.size = Pt(40)
            p_t.font.bold = True
            p_t.font.color.rgb = TEXT_TITLE

            # Subtitle
            sub_box = slide.shapes.add_textbox(Inches(2.0), Inches(3.7), Inches(9.3), Inches(0.8))
            p_s = sub_box.text_frame.paragraphs[0]
            p_s.text = item["subtitle"]
            p_s.font.size = Pt(22)
            p_s.font.bold = True
            p_s.font.color.rgb = TEXT_SUB

            # Meta Line
            meta_box = slide.shapes.add_textbox(Inches(2.0), Inches(4.7), Inches(9.3), Inches(0.6))
            p_m = meta_box.text_frame.paragraphs[0]
            p_m.text = item["meta"]
            p_m.font.size = Pt(14)
            p_m.font.color.rgb = TEXT_BODY

        else:
            # Content Slide
            # Category Pill
            cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(2.4), Inches(0.45))
            cat_box.fill.solid()
            cat_box.fill.fore_color.rgb = item["category_color"]
            cat_box.line.color.rgb = RGBColor(255, 255, 255)
            cat_box.line.width = Pt(1.5)
            p_cat = cat_box.text_frame.paragraphs[0]
            p_cat.text = f"● {item['category']}"
            p_cat.font.size = Pt(13)
            p_cat.font.bold = True
            p_cat.font.color.rgb = TEXT_WHITE
            p_cat.alignment = PP_ALIGN.CENTER

            # Slide counter
            num_box = slide.shapes.add_textbox(Inches(11.0), Inches(0.45), Inches(1.5), Inches(0.45))
            p_num = num_box.text_frame.paragraphs[0]
            p_num.text = item["slide_num"]
            p_num.font.size = Pt(14)
            p_num.font.bold = True
            p_num.font.color.rgb = TEXT_TITLE
            p_num.alignment = PP_ALIGN.RIGHT

            # Slide Title
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(11.7), Inches(0.8))
            p_t = title_box.text_frame.paragraphs[0]
            p_t.text = item["title"]
            p_t.font.size = Pt(28)
            p_t.font.bold = True
            p_t.font.color.rgb = TEXT_TITLE

            # Takeaway Line
            take_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.4))
            p_tk = take_box.text_frame.paragraphs[0]
            p_tk.text = f"Key Point: {item['takeaway']}"
            p_tk.font.size = Pt(15)
            p_tk.font.bold = True
            p_tk.font.color.rgb = TEXT_SUB

            # Main Content Card (Left)
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.25), Inches(8.5), Inches(4.7))
            card.fill.solid()
            card.fill.fore_color.rgb = CARD_BG
            card.line.color.rgb = BORDER_CYAN
            card.line.width = Pt(1.5)

            bullet_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(8.1), Inches(4.4))
            tf_b = bullet_box.text_frame
            tf_b.word_wrap = True

            for idx, (head, desc) in enumerate(item["bullets"]):
                p_head = tf_b.paragraphs[0] if idx == 0 else tf_b.add_paragraph()
                p_head.text = f"• {head}"
                p_head.font.size = Pt(16)
                p_head.font.bold = True
                p_head.font.color.rgb = TEXT_HEAD
                p_head.space_before = Pt(8) if idx > 0 else Pt(0)

                p_desc = tf_b.add_paragraph()
                p_desc.text = f"   {desc}"
                p_desc.font.size = Pt(14)
                p_desc.font.color.rgb = TEXT_BODY
                p_desc.space_after = Pt(10)

            # Metric Card (Right)
            kpi_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.55), Inches(2.25), Inches(2.95), Inches(4.7))
            kpi_card.fill.solid()
            kpi_card.fill.fore_color.rgb = RGBColor(240, 249, 255)
            kpi_card.line.color.rgb = item["category_color"]
            kpi_card.line.width = Pt(2)

            kpi_box = slide.shapes.add_textbox(Inches(9.7), Inches(3.2), Inches(2.65), Inches(2.8))
            tf_k = kpi_box.text_frame
            tf_k.word_wrap = True

            p_k1 = tf_k.paragraphs[0]
            p_k1.text = "KEY FACT"
            p_k1.font.size = Pt(12)
            p_k1.font.bold = True
            p_k1.font.color.rgb = TEXT_SUB
            p_k1.alignment = PP_ALIGN.CENTER

            p_k2 = tf_k.add_paragraph()
            p_k2.text = item["metric_box"][0]
            p_k2.font.size = Pt(22)
            p_k2.font.bold = True
            p_k2.font.color.rgb = item["category_color"]
            p_k2.alignment = PP_ALIGN.CENTER
            p_k2.space_before = Pt(8)

            p_k3 = tf_k.add_paragraph()
            p_k3.text = item["metric_box"][1]
            p_k3.font.size = Pt(12)
            p_k3.font.color.rgb = TEXT_BODY
            p_k3.alignment = PP_ALIGN.CENTER
            p_k3.space_before = Pt(10)

    prs.save(output_path)
    print(f"Bright Frutiger Aqua presentation saved to: {output_path}")

if __name__ == "__main__":
    output_pptx = os.path.abspath("WQS_SWOT_Frutiger_Aqua.pptx")
    create_bright_aqua_presentation(output_pptx)
